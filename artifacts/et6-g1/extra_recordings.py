import urllib.request,json,pathlib,hashlib,datetime,re,io,csv,concurrent.futures
from openpyxl import load_workbook
from fetch_index import fetch
OUT=pathlib.Path(__file__).resolve().parent
TASKS=[('AL-primary-gop','https://www.sos.alabama.gov/sites/default/files/05-29-2026/GOP%20Results.xlsx'),('IL-primary-senate','https://www.elections.il.gov/Downloads/ElectionOperations/ElectionResults/ByOffice/69/69-150-UNITED%20STATES%20SENATOR-2026GP.csv'),('IL-primary-governor','https://www.elections.il.gov/Downloads/ElectionOperations/ElectionResults/ByOffice/69/69-180-GOVERNOR%20AND%20LIEUTENANT%20GOVERNOR-2026GP.csv'),('HI-primary-summary','https://elections.hawaii.gov/wp-content/results/2026%20Primary/histatewide.pdf')]
def run(task):
    name,url=task
    try:
        with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'ElectionTracker-official-ENR-source-review/1.0'}),timeout=35) as r:body=r.read(); content=r.headers.get('Content-Type')
        receipt={'name':name,'url':url,'httpStatus':200,'sha256':hashlib.sha256(body).hexdigest(),'contentType':content,'byteLength':len(body),'checkedAt':datetime.datetime.now(datetime.timezone.utc).isoformat()}
        if url.endswith('.xlsx'):
            w=load_workbook(io.BytesIO(body),read_only=True,data_only=True);sheets={s.title:[list(v) for v in s.iter_rows(values_only=True)] for s in w};assert not re.search(r'\b(?:phone|email|street.address)\b',json.dumps(sheets,default=str),re.I);receipt['schema']={k:{'rows':len(v),'firstRows':v[:4]} for k,v in sheets.items()};(OUT/(name+'.xlsx')).write_bytes(body)
        elif url.endswith('.csv'):
            lines=list(csv.reader(io.StringIO(body.decode('utf-8-sig','replace'))));receipt['rows']=len(lines);receipt['firstRows']=lines[:4];assert not re.search(r'\b(?:phone|email|address)\b',str(lines[:4]),re.I);(OUT/(name+'.csv')).write_bytes(body)
        else:
            from pypdf import PdfReader
            pdf=PdfReader(io.BytesIO(body));text='\n'.join(p.extract_text() or '' for p in pdf.pages);assert not re.search(r'\b(?:phone|email)\b|[\w.-]+@[\w.-]+',text,re.I);receipt['pages']=len(pdf.pages);receipt['firstText']=text[:800];(OUT/(name+'.pdf')).write_bytes(body);(OUT/(name+'-text.txt')).write_text(text,encoding='utf-8')
        receipt['contactsRetained']=False
    except Exception as e:receipt={'name':name,'url':url,'httpStatus':getattr(e,'code',None),'error':str(e),'outcome':'transport_blocked'}
    (OUT/(name+'-receipt.json')).write_text(json.dumps(receipt,indent=2,ensure_ascii=False,default=str)+'\n',encoding='utf-8');return receipt
with concurrent.futures.ThreadPoolExecutor(max_workers=4) as p:
    for r in p.map(run,TASKS):print(json.dumps(r,default=str))
